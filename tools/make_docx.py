#!/usr/bin/env python3
"""Erzeugt die leere Word-Vorlage des Prüfprotokolls (A4 quer, gleiche Struktur wie das PDF).

Aufruf: python3 make_docx.py Ausgabe.docx
"""
import sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY, BLUE, GREY, INK = '173660', '2F75B5', '667085', '1F2937'
LINE, SOFT, SOFTLINE = 'B8C2CC', 'F4F7FA', 'D9E2F3'
OK, WARN, BAD = '2E7D32', 'C77700', 'B42318'
FONT = 'Arial'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

CHECKPOINTS = [
    'Treppenstufen vollständig und sicher begehbar', 'Stufen ohne Beschädigungen oder Stolperstellen',
    'Treppenlauf standsicher, keine erkennbaren Verformungen', 'Podeste sicher und frei von Schäden',
    'Handläufe fest und vollständig', 'Geländer vollständig und standsicher',
    'Befestigungen und Verbindungselemente augenscheinlich intakt', 'Schweißnähte ohne erkennbare Schäden',
    'Stahlbauteile ohne erhebliche Korrosion', 'Treppe frei von Hindernissen und Gegenständen',
    'Flucht- und Rettungsweg jederzeit nutzbar', 'Bodenbeläge / Gitterroste trittsicher',
    'Keine losen oder fehlenden Bauteile', 'Beleuchtung ausreichend und funktionsfähig',
    'Sicherheits- und Fluchtwegkennzeichnung vorhanden', 'Türen im Fluchtweg frei zugänglich und funktionsfähig',
    'Keine erkennbaren Gefährdungen durch Witterungseinflüsse', 'Entwässerung / Wasserablauf funktionsfähig',
    'Auftritts- und Stufenkanten augenscheinlich ordnungsgemäß', 'Geländerfüllungen / Absturzsicherung vollständig',
    'Korrosionsschutz / Beschichtung augenscheinlich intakt', 'Anschlüsse an Baukörper augenscheinlich intakt',
    'Fluchtwegbreite augenscheinlich nicht eingeschränkt', 'Zugang zur Treppe jederzeit möglich',
]
GRUNDLAGEN = [
    'DIN 18065 - Gebäudetreppen', 'DIN EN 1090 - Ausführung von Stahltragwerken',
    'Arbeitsstättenverordnung (ArbStättV)', 'ASR A2.3 - Fluchtwege und Notausgänge',
    'einschlägige DGUV-Vorschriften und DGUV-Regeln', 'objektspezifische Anforderungen',
    'Herstellerangaben und technische Dokumentationen', 'bauordnungsrechtliche Vorgaben / genehmigte Planunterlagen',
]
ERGEBNIS = [
    ('Keine Mängel festgestellt', 'Die geprüfte Treppe weist im Rahmen der durchgeführten Sicht- und Funktionsprüfung keine erkennbaren Mängel auf.', OK),
    ('Geringfügige Mängel festgestellt', 'Die Nutzung ist nach fachlicher Bewertung weiterhin möglich. Eine zeitnahe Mängelbeseitigung ist erforderlich.', WARN),
    ('Erhebliche Mängel festgestellt', 'Eine kurzfristige Mängelbeseitigung ist erforderlich. Die weitere Nutzung ist unter Berücksichtigung der konkreten Gefährdung zu bewerten.', BAD),
    ('Sicherheitsrelevanter Mangel festgestellt', 'Die Nutzung ist bis zur Beseitigung bzw. geeigneten Absicherung des Mangels einzuschränken oder die Treppe zu sperren.', BAD),
    ('Weitergehende fachkundige Prüfung erforderlich', 'Eine abschließende Beurteilung ist durch eine entsprechend qualifizierte Fachperson / einen Sachverständigen erforderlich.', BLUE),
]
PRUEFHINWEIS = 'Die Prüfung ist objektbezogen durchzuführen. Anzuwendende Vorschriften, genehmigte Planunterlagen, Herstellerangaben und weitere technische Grundlagen sind vor Ort festzulegen und bei Bedarf zu dokumentieren.'
GRUNDLAGEN_HINWEIS = 'Die tatsächliche Auswahl und Rangfolge der Prüfgrundlagen ist objektspezifisch festzulegen. Die Checkfelder dienen der Dokumentation und ersetzen keine rechtliche oder fachkundige Prüfung.'
CHECKLISTE_HINWEIS = 'Bei Mängeln bitte fortlaufend in Abschnitt 4 eintragen und ggf. Foto in Abschnitt 5 zuordnen.'
BEWERTUNGSHINWEIS = 'Die Priorität ist unter Berücksichtigung der konkreten Gefährdung, der Nutzung der Treppe und der betrieblichen Rahmenbedingungen festzulegen. Bei sicherheitsrelevanten Mängeln ist eine geeignete Absicherung bzw. Sperrung zu bewerten.'
DOKU_HINWEIS = 'Mit der Unterschrift wird die Durchführung und Dokumentation der Prüfung bestätigt. Technische oder rechtliche Schlussfolgerungen sind auf die im Prüfauftrag vorgesehenen Kompetenzen zu beschränken.'
FOOTER = 'DHRW Engineering GmbH · Edisonallee 39 · 89231 Neu-Ulm · Tel. 0731 14113924 · info@dhrw.group'

CONTENT_W = 27.7  # cm


def rgb(h):
    return RGBColor.from_string(h)


def style_run(run, size=9, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return run


def fmt(p, before=0, after=0, align=None, keep_next=False, line=None):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if align is not None:
        p.alignment = align
    if keep_next:
        pf.keep_with_next = True
    if line:
        pf.line_spacing = line
    return p


def text(container, s='', size=9, bold=False, color=INK, align=None, before=0, after=0, reuse=False, italic=False):
    """Absatz in Zelle/Dokument. reuse=True nimmt den ersten (leeren) Absatz einer Zelle."""
    p = container.paragraphs[0] if reuse else container.add_paragraph()
    if s:
        style_run(p.add_run(s), size, bold, color, italic)
    return fmt(p, before, after, align)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_)
    tcPr.append(shd)


def borders(cell, color=LINE, sz=4, val='single', sides=('top', 'left', 'bottom', 'right')):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders')
        tcPr.append(b)
    for side in ('top', 'left', 'bottom', 'right'):
        el = b.find(qn(f'w:{side}'))
        if el is None:
            el = OxmlElement(f'w:{side}')
            b.append(el)
        if side in sides:
            el.set(qn('w:val'), val); el.set(qn('w:sz'), str(sz)); el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)
        else:
            el.set(qn('w:val'), 'nil')


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)


def valign(cell, v='center'):
    tcPr = cell._tc.get_or_add_tcPr()
    e = OxmlElement('w:vAlign'); e.set(qn('w:val'), v); tcPr.append(e)


def row_height(row, cm, rule='atLeast'):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:trHeight')
    h.set(qn('w:val'), str(int(cm * 567))); h.set(qn('w:hRule'), rule)
    trPr.append(h)
    cs = OxmlElement('w:cantSplit')
    trPr.append(cs)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))


def make_table(container, rows, widths, indent_cm=0):
    t = container.add_table(rows=rows, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = tblPr.find(qn('w:tblW'))
    if tw is None:
        tw = OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:w'), str(int(sum(widths) * 567))); tw.set(qn('w:type'), 'dxa')
    if indent_cm:
        ind = OxmlElement('w:tblInd'); ind.set(qn('w:w'), str(int(indent_cm * 567))); ind.set(qn('w:type'), 'dxa'); tblPr.append(ind)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w * 567)))
    for r in t.rows:
        for c, w in zip(r.cells, widths):
            c.width = Cm(w)
            borders(c, sides=())  # standardmäßig ohne Rahmen
    return t


def checkbox(p, size=11):
    """Anklickbares Kontrollkästchen (Word-Inhaltssteuerelement)."""
    sdt = parse_xml(
        f'<w:sdt xmlns:w="{W_NS}" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
        '<w:sdtPr><w14:checkbox><w14:checked w14:val="0"/>'
        '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/><w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/>'
        '</w14:checkbox></w:sdtPr><w:sdtContent><w:r><w:rPr>'
        '<w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic" w:hint="eastAsia"/>'
        f'<w:color w:val="{BLUE}"/><w:sz w:val="{size * 2}"/></w:rPr><w:t>☐</w:t></w:r></w:sdtContent></w:sdt>')
    p._p.append(sdt)


def option(p, label, size=9, gap=''):
    checkbox(p)
    style_run(p.add_run(f' {label}{gap}'), size)


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    fmt(p)


def heading(doc, eyebrow, title, first=False):
    if not first:
        page_break(doc)
    text(doc, eyebrow, 8, True, BLUE, before=0, after=2).paragraph_format.keep_with_next = True
    p = text(doc, title, 17, True, NAVY, after=8)
    p.paragraph_format.keep_with_next = True


def field_row(doc, items, widths, min_cm=1.15, gap_after=6):
    """Reihe beschrifteter Felder: items = [Beschriftung, ...]; Lücken zwischen den Feldern als schmale Spalten."""
    cols, ws = [], []
    for i, w in enumerate(widths):
        if i:
            ws.append(0.4)
        ws.append(w)
    t = make_table(doc, 1, ws)
    row_height(t.rows[0], min_cm)
    ci = 0
    for i, c in enumerate(t.rows[0].cells):
        if i % 2 == 0:
            borders(c)
            cell_margins(c, 50, 50, 110, 110)
            text(c, items[ci], 8, True, GREY, reuse=True)
            text(c, '', 10)
            ci += 1
    spacer(doc, gap_after)
    return t


def spacer(doc, pt=6):
    p = doc.add_paragraph()
    fmt(p, 0, 0)
    p.paragraph_format.line_spacing = Pt(pt)
    r = p.add_run('')
    r.font.size = Pt(2)
    return p


def note_box(doc, title, body, width=CONTENT_W):
    t = make_table(doc, 1, [width])
    c = t.rows[0].cells[0]
    shade(c, SOFT); borders(c, SOFTLINE)
    cell_margins(c, 100, 100, 200, 200)
    text(c, title, 9.5, True, NAVY, reuse=True, after=3)
    text(c, body, 8.5, False, INK)
    spacer(doc, 6)


def text_box(doc, label, min_cm, label_style=None):
    if label:
        ls = label_style or (8.3, True, GREY)
        text(doc, label, ls[0], ls[1], ls[2], after=2).paragraph_format.keep_with_next = True
    t = make_table(doc, 1, [CONTENT_W])
    c = t.rows[0].cells[0]
    borders(c); cell_margins(c, 80, 80, 140, 140)
    row_height(t.rows[0], min_cm)
    text(c, '', 10, reuse=True)
    spacer(doc, 8)


def add_field(run_parent, instr):
    for kind, txt in (('begin', None), (None, instr), ('end', None)):
        r = run_parent.add_run()
        style_run(r, 7, False, GREY)
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' {txt} '; r._r.append(it)


TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
              'textDirection', 'tcFitText', 'vAlign', 'hideMark']
TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW',
               'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook']


def reorder(el, order):
    """Kinder in die vom Word-Schema verlangte Reihenfolge bringen."""
    kids = list(el)
    key = lambda k: order.index(k.tag.split('}')[1]) if k.tag.split('}')[1] in order else len(order)
    for k in kids:
        el.remove(k)
    for k in sorted(kids, key=key):
        el.append(k)


def normalize(doc):
    roots = [doc.element.body, doc.sections[0].header._element, doc.sections[0].footer._element]
    for root in roots:
        for el in root.iter(qn('w:tcPr')):
            reorder(el, TCPR_ORDER)
        for el in root.iter(qn('w:tblPr')):
            reorder(el, TBLPR_ORDER)


def build(path):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.font.size = Pt(9)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.5)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0.6)

    # Kopfzeile: marineblaues Band über die ganze Seitenbreite
    hdr = sec.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    ht = hdr.add_table(1, 2, Cm(29.7))
    ht.autofit = False
    tblPr = ht._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    ind = OxmlElement('w:tblInd'); ind.set(qn('w:w'), str(-567)); ind.set(qn('w:type'), 'dxa'); tblPr.append(ind)
    for gc, w in zip(ht._tbl.tblGrid.findall(qn('w:gridCol')), (14.85, 14.85)):
        gc.set(qn('w:w'), str(int(w * 567)))
    row_height(ht.rows[0], 1.15, 'exact')
    for i, c in enumerate(ht.rows[0].cells):
        c.width = Cm(14.85)
        shade(c, NAVY); borders(c, sides=())
        cell_margins(c, 0, 0, 567, 567)
        valign(c)
        if i == 0:
            text(c, 'DHRW ENGINEERING GMBH', 12, True, 'FFFFFF', reuse=True)
        else:
            text(c, 'PRÜFPROTOKOLL', 11, True, 'FFFFFF', align=WD_ALIGN_PARAGRAPH.RIGHT, reuse=True)
    hp._p.getparent().remove(hp._p)
    hdr._element.append(OxmlElement('w:p'))

    # Fußzeile
    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles['Normal']
    pPr = fp._p.get_or_add_pPr()
    pb = OxmlElement('w:pBdr')
    top = OxmlElement('w:top'); top.set(qn('w:val'), 'single'); top.set(qn('w:sz'), '4'); top.set(qn('w:space'), '4'); top.set(qn('w:color'), LINE)
    pb.append(top); pPr.append(pb)
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab'); tab.set(qn('w:val'), 'right'); tab.set(qn('w:pos'), str(int(CONTENT_W * 567))); tabs.append(tab); pPr.append(tabs)
    style_run(fp.add_run(FOOTER + '\tSeite '), 7, False, GREY)
    add_field(fp, 'PAGE')
    style_run(fp.add_run(' von '), 7, False, GREY)
    add_field(fp, 'NUMPAGES')

    # ---- 1 · Allgemeine Angaben ----
    heading(doc, '1 · ALLGEMEINE ANGABEN', 'Objekt- und Prüfungsdaten', first=True)
    field_row(doc, ['Dokument-Nr.', 'Revision', 'Prüfdatum', 'Prüfer'], [8.0, 4.6, 6.3, 7.6])
    field_row(doc, ['Objekt', 'Gebäude / Bereich'], [13.65, 13.65])
    field_row(doc, ['Standort', 'Treppen-Nr.'], [13.65, 13.65])

    t = make_table(doc, 2, [13.65, 0.4, 13.65])
    for ci, (name, opts) in ((0, ('Treppentyp', ['Fluchttreppe', 'Rettungswegtreppe', 'Sicherheitstreppe'])),
                             (2, ('Bauart', ['Stahl', 'Stahlbeton', 'Sonstige:']))):
        c = t.rows[0].cells[ci]
        shade(c, SOFT); cell_margins(c, 50, 50, 140, 140)
        text(c, name, 8, True, NAVY, reuse=True)
        c2 = t.rows[1].cells[ci]
        cell_margins(c2, 90, 40, 140, 40)
        p = c2.paragraphs[0]
        for o in opts:
            option(p, o, 9, gap='   ')
        if name == 'Bauart':
            style_run(p.add_run('_______________'), 9, False, LINE)
    spacer(doc, 8)
    text_box(doc, 'Anwesende Personen', 1.3)
    note_box(doc, 'Prüfhinweis', PRUEFHINWEIS)

    # ---- 2 · Prüfgrundlagen ----
    heading(doc, '2 · PRÜFGRUNDLAGEN', 'Prüfgrundlagen und Anwendungsbereich')
    t = make_table(doc, len(GRUNDLAGEN), [CONTENT_W])
    for r, g in zip(t.rows, GRUNDLAGEN):
        c = r.cells[0]
        borders(c); cell_margins(c, 70, 70, 200, 100)
        p = c.paragraphs[0]
        option(p, g, 9.5)
    spacer(doc, 10)
    text_box(doc, 'Weitere Grundlagen / objektspezifische Festlegungen', 2.0)
    note_box(doc, 'Hinweis', GRUNDLAGEN_HINWEIS)

    # ---- 3 · Sicht- und Funktionsprüfung ----
    heading(doc, '3 · SICHT- UND FUNKTIONSPRÜFUNG', 'Prüfcheckliste')
    text(doc, 'i. O. = in Ordnung  ·  n. a. = nicht anwendbar', 8.5, False, GREY, after=5)
    widths = [1.3, 13.3, 2.0, 2.0, 2.0, 7.1]
    t = make_table(doc, 1 + len(CHECKPOINTS), widths)
    for c, h_ in zip(t.rows[0].cells, ['Nr.', 'Prüfpunkt', 'i. O.', 'Mangel', 'n. a.', 'Bemerkung']):
        shade(c, SOFT); borders(c); cell_margins(c, 70, 70, 100, 100)
        text(c, h_, 8, True, GREY, reuse=True,
             align=WD_ALIGN_PARAGRAPH.CENTER if h_ in ('Nr.', 'i. O.', 'Mangel', 'n. a.') else None)
    repeat_header(t.rows[0])
    for i, (r, name) in enumerate(zip(t.rows[1:], CHECKPOINTS), start=1):
        row_height(r, 0.9)
        for ci, c in enumerate(r.cells):
            borders(c); valign(c); cell_margins(c, 40, 40, 100, 100)
        text(r.cells[0], str(i), 8.5, True, NAVY, reuse=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        text(r.cells[1], name, 8.6, reuse=True)
        for ci in (2, 3, 4):
            p = text(r.cells[ci], '', reuse=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            checkbox(p, 12)
        text(r.cells[5], '', 8.2, reuse=True)
    spacer(doc, 8)
    note_box(doc, 'Dokumentation', CHECKLISTE_HINWEIS)

    # ---- 4 · Festgestellte Mängel ----
    heading(doc, '4 · FESTGESTELLTE MÄNGEL', 'Mängelerfassung')
    widths = [1.2, 7.2, 2.6, 7.2, 5.4, 4.1]
    t = make_table(doc, 6, widths)
    for c, h_ in zip(t.rows[0].cells, ['Nr.', 'Feststellung / Mangel', 'Priorität', 'Erforderliche Maßnahme', 'Verantwortlich', 'Termin']):
        shade(c, SOFT); borders(c); cell_margins(c, 70, 70, 100, 100)
        text(c, h_, 8, True, GREY, reuse=True, align=WD_ALIGN_PARAGRAPH.CENTER if h_ == 'Nr.' else None)
    repeat_header(t.rows[0])
    for k, r in enumerate(t.rows[1:], start=1):
        row_height(r, 2.1)
        for c in r.cells:
            borders(c); cell_margins(c, 80, 60, 100, 100)
        text(r.cells[0], str(k), 8.5, True, NAVY, reuse=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        c = r.cells[2]
        first = True
        for lbl in ('hoch', 'mittel', 'niedrig'):
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            option(p, lbl, 8.2)
            fmt(p, 0, 1)
        for ci in (1, 3, 4, 5):
            text(r.cells[ci], '', 8.6, reuse=True)
    text(doc, 'Weitere Zeilen: in der letzten Zelle die Tabulatortaste drücken.', 8, False, GREY, before=3, after=6, italic=True)
    note_box(doc, 'Bewertungshinweis', BEWERTUNGSHINWEIS)

    # ---- 5 · Fotodokumentation ----
    heading(doc, '5 · FOTODOKUMENTATION', 'Fotodokumentation und Anlagen')
    widths = [13.65, 0.4, 13.65]
    for pair in range(3):
        if pair == 2:
            page_break(doc)
            text(doc, '5 · FOTODOKUMENTATION', 8, True, BLUE, after=2)
            text(doc, 'Fotodokumentation und Anlagen (Fortsetzung)', 17, True, NAVY, after=8)
        t = make_table(doc, 3, widths)
        for ci in (0, 2):
            n = pair * 2 + (1 if ci == 2 else 0) + 1
            c = t.rows[0].cells[ci]
            borders(c, sides=('top', 'left', 'right')); cell_margins(c, 80, 20, 140, 140)
            text(c, f'Foto {n}', 9, True, NAVY, reuse=True)
            c = t.rows[1].cells[ci]
            borders(c, color=BLUE, sz=4, val='dashed', sides=('top', 'left', 'bottom', 'right')); valign(c)
            shade(c, SOFT)
            text(c, 'Foto hier einfügen', 9, True, GREY, reuse=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            c = t.rows[2].cells[ci]
            borders(c, sides=('bottom', 'left', 'right')); cell_margins(c, 60, 80, 140, 140)
            text(c, 'Bemerkung:', 8.3, True, GREY, reuse=True)
        row_height(t.rows[1], 5.2, 'exact')
        row_height(t.rows[2], 1.2)
        spacer(doc, 10)

    # ---- 6 · Ergebnis ----
    heading(doc, '6 · ERGEBNIS DER PRÜFUNG', 'Ergebnis und fachliche Bewertung')
    t = make_table(doc, len(ERGEBNIS), [0.3, 1.2, 26.2])
    for r, (title, desc, colr) in zip(t.rows, ERGEBNIS):
        row_height(r, 1.5)
        shade(r.cells[0], colr)
        for c in r.cells[1:]:
            valign(c); cell_margins(c, 80, 80, 120, 120)
        borders(r.cells[0], color=colr)
        borders(r.cells[1], sides=('top', 'bottom')); borders(r.cells[2], sides=('top', 'bottom', 'right'))
        p = text(r.cells[1], '', reuse=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        checkbox(p, 14)
        text(r.cells[2], title, 10, True, NAVY, reuse=True, after=1)
        text(r.cells[2], desc, 8.6)
        borders(r.cells[0], color=colr)
    spacer(doc, 10)
    text_box(doc, 'Zusatzbewertung / Begründung', 2.2)

    # ---- 7-8 · Bemerkungen ----
    heading(doc, '7-8 · BEMERKUNGEN UND MASSNAHMEN', 'Bemerkungen, Empfehlungen und Nachprüfung')
    text_box(doc, '7. Bemerkungen', 3.6, (10, True, NAVY))
    text_box(doc, '8. Empfehlungen / weitere Maßnahmen', 3.6, (10, True, NAVY))
    text(doc, 'Erforderliche Nachprüfung', 10, True, NAVY, after=4)
    p = doc.add_paragraph()
    option(p, 'nein', 9.5, gap='          ')
    option(p, 'ja', 9.5, gap='          ')
    style_run(p.add_run('Termin: '), 9.5, True, GREY)
    style_run(p.add_run('____________________'), 9.5, False, LINE)

    # ---- 9 · Abschluss ----
    heading(doc, '9 · ABSCHLUSS', 'Abschluss und nächste Prüfung')
    t = make_table(doc, 7, [13.65, 0.4, 13.65])
    plan = [('title', None), ('label', 'Name'), ('box', 0.9), ('label', 'Unterschrift'), ('box', 2.4), ('label', 'Datum'), ('box', 0.9)]
    for r, (kind, arg) in zip(t.rows, plan):
        for ci, title in ((0, 'Prüfer'), (2, 'Verantwortlicher Betreiber')):
            c = r.cells[ci]
            if kind == 'title':
                text(c, title, 11, True, NAVY, reuse=True, before=4, after=4)
            elif kind == 'label':
                text(c, arg, 8.3, True, GREY, reuse=True, before=4, after=1)
            else:
                borders(c); cell_margins(c, 60, 60, 110, 110)
                row_height(r, arg)
                text(c, '', 10, reuse=True)
    spacer(doc, 12)
    t = make_table(doc, 2, [6.8, 0.6, 20.3])
    text(t.rows[0].cells[0], 'Nächster Prüftermin', 11, True, NAVY, reuse=True, after=4)
    c = t.rows[1].cells[0]
    borders(c); cell_margins(c, 60, 60, 110, 110)
    row_height(t.rows[1], 0.9)
    text(c, '', 10, reuse=True)
    c = t.rows[0].cells[2]
    text(c, 'Dokumentationshinweis:', 8.3, True, GREY, reuse=True, after=3)
    text(c, DOKU_HINWEIS, 8.3)

    # Dokumenteigenschaften
    cp = doc.core_properties
    cp.title = 'Prüfprotokoll (leere Vorlage)'
    cp.author = 'DHRW Engineering GmbH'
    normalize(doc)
    doc.save(path)


if __name__ == '__main__':
    build(sys.argv[1] if len(sys.argv) > 1 else 'Leeres_Pruefprotokoll.docx')
